"""PDF (reportlab) and Excel (openpyxl) exporters for payslips and tabular reports."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND = colors.HexColor("#4f46e5")


def _money(v) -> str:
    return f"Rs. {v:,.2f}"


def table_pdf(title: str, columns: list[tuple[str, str]], rows: list[dict], subtitle: str = "") -> bytes:
    buf = io.BytesIO()
    wide = len(columns) > 6
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if wide else A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=title)
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    head = styles["BodyText"].clone("head", fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(f"<font color='#4f46e5'>WorkForce AI Pro</font> — {title}", styles["Title"]),
             Paragraph(f"{subtitle} Generated {datetime.now():%d %b %Y %H:%M} · {len(rows)} record(s)", styles["Normal"]),
             Spacer(1, 6 * mm)]
    data = [[Paragraph(label, head) for _, label in columns]]
    for r in rows:
        data.append([Paragraph("" if r.get(k) is None else (f"{r[k]:,.2f}" if isinstance(r.get(k), float) else str(r[k])), cell)
                     for k, _ in columns])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f4fb")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d4d4e4")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(t)
    doc.build(story)
    return buf.getvalue()


def table_xlsx(title: str, columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([label for _, label in columns])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F46E5")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([r.get(k) for k, _ in columns])
    for i, (k, label) in enumerate(columns, start=1):
        width = max([len(str(label))] + [len(str(r.get(k) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(45, width + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def payslip_pdf(p: dict, emp: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=f"Payslip {p['period']}")
    styles = getSampleStyleSheet()
    story = [Paragraph("<font color='#4f46e5'>WorkForce AI Pro</font>", styles["Title"]),
             Paragraph(f"Payslip for {datetime.strptime(p['period'], '%Y-%m'):%B %Y}", styles["Heading2"]),
             Spacer(1, 4 * mm)]
    info = [["Employee", emp["full_name"], "Employee Code", emp["employee_code"]],
            ["Department", emp.get("department") or "-", "Designation", emp.get("job_title") or "-"],
            ["Status", p["status"], "Generated", p["generated_at"][:10] if p.get("generated_at") else "-"]]
    t = Table(info, colWidths=[30 * mm, 55 * mm, 32 * mm, 55 * mm])
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d4d4e4")),
                           ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8f8fd"))]))
    story += [t, Spacer(1, 6 * mm)]
    earnings = [("Basic", p["basic"]), ("House Rent Allowance", p["hra"]), ("Special Allowance", p["allowances"]),
                ("Overtime Pay", p["overtime_pay"]), ("Bonus", p["bonus"]), ("Incentives", p["incentives"])]
    deductions = [("Income Tax (TDS)", p["tax"]), ("Provident Fund", p["pf"]),
                  ("Unpaid Leave", p["unpaid_leave_deduction"]), ("", None), ("", None), ("", None)]
    rows = [["Earnings", "Amount", "Deductions", "Amount"]]
    for (el, ev), (dl, dv) in zip(earnings, deductions):
        rows.append([el, _money(ev), dl, _money(dv) if dv is not None else ""])
    rows.append(["Gross Earnings", _money(p["gross"]), "Total Deductions", _money(p["total_deductions"])])
    t = Table(rows, colWidths=[50 * mm, 36 * mm, 50 * mm, 36 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d4d4e4")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#eef0ff")),
    ]))
    story += [t, Spacer(1, 6 * mm)]
    net = Table([["NET PAY", _money(p["net"])]], colWidths=[86 * mm, 86 * mm])
    net.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BRAND), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                             ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 13),
                             ("ALIGN", (1, 0), (1, 0), "RIGHT"), ("TOPPADDING", (0, 0), (-1, -1), 8),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [net, Spacer(1, 8 * mm),
              Paragraph("This is a system-generated payslip and does not require a signature.", styles["Italic"])]
    doc.build(story)
    return buf.getvalue()
